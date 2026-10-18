"""Reading an uploaded budget file (XLSX or CSV) and applying it to budgets."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from datetime import date
from typing import Any

from sqlalchemy.orm import Session

from fiscus.domain import budget_import as dom
from fiscus.domain.uploads import ImportError_, decode, read_rows
from fiscus.models import Budget, Category, User
from fiscus.money import round_div
from fiscus.services import budgeting
from fiscus.services.access import owned
from fiscus.services.ledger import category_named

MAX_UNZIPPED = 50 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLS = 60


@dataclass
class Sheet:
    name: str
    rows: list[list[Any]]


def read_sheets(data: bytes, filename: str) -> list[Sheet]:
    if data[:4] == b'PK\x03\x04' or filename.lower().endswith(('.xlsx', '.xlsm')):
        return _read_xlsx(data)
    return [Sheet('CSV', [list(r) for r in read_rows(decode(data))])]


def _read_xlsx(data: bytes) -> list[Sheet]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_UNZIPPED:
                raise ImportError_('That spreadsheet is too large to read')
    except zipfile.BadZipFile as exc:
        raise ImportError_('That is not an Excel (.xlsx) file') from exc
    # defusedxml is installed, so openpyxl parses with it: no entity expansion.
    import defusedxml  # noqa: F401
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_('That spreadsheet could not be opened') from exc
    sheets = []
    try:
        for ws in wb.worksheets:
            rows: list[list[Any]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    break
                rows.append(list(row[:MAX_COLS]))
            sheets.append(Sheet(ws.title, rows))
    finally:
        wb.close()
    return sheets


@dataclass
class Parsed:
    sheet: str
    sheets: list[str]
    lines: list[dom.Line]


def parse(data: bytes, filename: str, sheet: str | None, *, has_loan: bool) -> Parsed:
    sheets = read_sheets(data, filename)
    names = [s.name for s in sheets]
    order = [s for s in sheets if s.name == sheet] + [s for s in sheets if s.name != sheet]
    for s in order:
        header = dom.find_header(s.rows)
        if header is not None:
            return Parsed(s.name, names, dom.parse(s.rows, header, has_loan=has_loan))
    raise ImportError_(
        'Could not find a header row with Item, Amount and Frequency columns. '
        'Add those headings above your lines, or start from the template.'
    )


@dataclass(frozen=True)
class Chosen:
    category: str
    total_cents: int  # for one cycle, as typed
    cycle_months: int = 1

    @property
    def monthly_cents(self) -> int:
        return round_div(self.total_cents, self.cycle_months)


def apply(
    db: Session,
    user: User,
    chosen: list[Chosen],
    today: date,
    *,
    replace: bool,
    use_in_forecast: bool,
) -> tuple[int, int]:
    """Sum the chosen lines per category and set those budgets. Returns (categories,
    monthly total).

    An imported amount is the category's whole budget: its bills count
    towards it, and the rest is budgeted on top of them (D64). A category
    takes the longest cycle of its lines, so a yearly cost in it is saved up.
    A category holding a bill keeps a line, cleared or not.
    """
    lines: dict[int, tuple[Category, list[Chosen]]] = {}
    for c in chosen:
        cat = category_named(db, user, c.category)
        lines.setdefault(cat.id, (cat, []))[1].append(c)
    monthly_total = 0
    for cat, picked in lines.values():
        cycle = max(c.cycle_months for c in picked)
        if cat.is_income or cycle not in budgeting.CYCLES:
            cycle = 1
        if all(c.cycle_months == cycle for c in picked):
            total = sum(c.total_cents for c in picked)
        else:
            total = sum(c.monthly_cents for c in picked) * cycle
        budgeting.set_total(
            db,
            user,
            cat,
            total,
            today,
            use_in_forecast=use_in_forecast,
            add=not replace,
            cycle_months=cycle,
            whole=True,
        )
        monthly_total += sum(c.monthly_cents for c in picked)
    budgeting.follow_bills(db, user, today)
    return len(lines), monthly_total


def clear(db: Session, user: User) -> int:
    """Remove every budget line, for an import that starts clean. What was
    taken off by hand stays on record. Returns how many were removed."""
    rows = list(db.execute(owned(Budget, user)).scalars())
    for row in rows:
        db.delete(row)
    db.flush()
    return len(rows)


def template_csv() -> str:
    return (
        'Item,Amount,Frequency,Notes\n'
        'Groceries,240,Weekly,Two adults\n'
        'Internet,79,Monthly,\n'
        'Council rates,2960,Annual,\n'
        'Car rego,1109,Annual,\n'
    )
